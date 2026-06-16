# generate_reports.py
# Main builder script for generating unique internship reports for all students.
# Implements design themes, spintax parsing, and a two-pass page number extraction.

import os
import re
import csv
import sys
import shutil
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Import report data
from report_data import STUDENTS, THEMES, REPORT_CONTENT

# Directories
WORKSPACE_DIR = "/home/ihjas/Documents/GitHub/SSC-FILES/2026/INTERNSHIP"
OUTPUT_DIR = os.path.join(WORKSPACE_DIR, "REPORTS")
LOGO_PATH = os.path.join(WORKSPACE_DIR, "logo_extracted-000.png")

# Ensure output directory exists
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Helper function for cell background color
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Helper function to remove cell borders (for clean cover pages)
def remove_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

# Deterministic spintax choice
def get_deterministic_choice(student_index, placeholder_index, num_options):
    val = (student_index * 137 + placeholder_index * 79 + 503)
    val = (val ^ (val >> 8)) * 0x85ebca6b
    val = (val ^ (val >> 13)) * 0xc2b2ae35
    val = val ^ (val >> 16)
    return abs(val) % num_options

# Spintax parser
def parse_spintax(text, student_index):
    placeholder_index = 0
    def replace(match):
        nonlocal placeholder_index
        options = match.group(1).split('|')
        idx = get_deterministic_choice(student_index, placeholder_index, len(options))
        placeholder_index += 1
        return options[idx]
    
    pattern = re.compile(r'\{([^{}]+)\}')
    while pattern.search(text):
        text = pattern.sub(replace, text)
    return text

# Helper function to add standard paragraph
def add_p(doc, spintax, student_index, font_name, size_pt, color_rgb=None, bold=False, italic=False, space_after=6, alignment=None, line_spacing=1.15):
    text = parse_spintax(spintax, student_index)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    if alignment is not None:
        p.alignment = alignment
    
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)
    return p

# Helper function to add a list item
def add_list_item(doc, spintax, student_index, font_name, size_pt, bullet_char="• ", color_rgb=None, space_after=4, left_indent=0.25):
    text = parse_spintax(spintax, student_index)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(left_indent)
    p.paragraph_format.first_line_indent = Inches(-left_indent)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    run_bullet = p.add_run(bullet_char)
    run_bullet.font.name = font_name
    run_bullet.font.size = Pt(size_pt)
    run_bullet.font.bold = True
    if color_rgb:
        run_bullet.font.color.rgb = RGBColor(*color_rgb)
        
    run_text = p.add_run(text)
    run_text.font.name = font_name
    run_text.font.size = Pt(size_pt)
    return p

# Table Generator
def build_styled_table(doc, table_type, theme, student_index):
    font_name = theme["font"]
    primary_color_rgb = theme["primary_color"]
    primary_hex = "{:02X}{:02X}{:02X}".format(*primary_color_rgb)
    
    if table_type == "loss_parameters":
        # Table 3.1: Standard Loss Parameters Used in FTTH Networks
        data = [
            ("Fiber Path Loss @ 1310 nm / 1490 nm", "0.40 dB per km"),
            ("Fiber Path Loss @ 1550 nm", "0.25 dB per km"),
            ("Fusion Splice Junction", "0.10 dB per splice"),
            ("SC/UPC Physical Connector Pair", "0.35 dB per connector"),
            ("SC/APC Physical Connector Pair", "0.25 dB per connector"),
            ("1 × 2 Splitter Module", "3.0 dB average insertion loss"),
            ("1 × 16 Splitter Module", "13.0 dB average insertion loss"),
            ("1 × 32 Splitter Module", "16.0 dB average insertion loss")
        ]
        headers = ["Network Element Component", "Standard Loss Value"]
        caption = "Table 3.1: Standard Loss Parameters Used in FTTH Networks"
        
    elif table_type == "color_coding":
        # Table 6.1: TIA-598-C Color Coding Identification
        data = [
            ("1", "Blue", "BL", "7", "Red", "RD"),
            ("2", "Orange", "OR", "8", "Black", "BK"),
            ("3", "Green", "GR", "9", "Yellow", "YL"),
            ("4", "Brown", "BR", "10", "Violet", "VI"),
            ("5", "Slate (Grey)", "SL", "11", "Rose (Pink)", "RS"),
            ("6", "White", "WH", "12", "Aqua", "AQ")
        ]
        headers = ["No.", "Fiber Color", "Symbol", "No.", "Fiber Color", "Symbol"]
        caption = "Table 6.1: TIA-598-C Color Coding Identification"
        
    elif table_type == "splicing_comparison":
        # Table 6.2: Comparison of Fusion and Mechanical Splicing
        data = [
            ("Mechanism", 
             "Uses a high-voltage electric arc to melt and fuse glass tips permanently.", 
             "Physically aligns two fiber ends inside a sleeve using index-matching gel."),
            ("Signal Loss", 
             "Extremely low (< 0.05 dB to 0.1 dB).", 
             "Slightly higher (0.2 dB to 0.5 dB)."),
            ("Pros & Cons", 
             "Most reliable, but requires an expensive fusion splicer machine.", 
             "Fast setup, low initial tool cost. Ideal for home drop termination.")
        ]
        headers = ["Feature", "Fusion Splicing", "Mechanical / Fast Connectors"]
        caption = "Table 6.2: Comparison of Fusion and Mechanical Splicing"
        
    elif table_type == "troubleshooting_matrix":
        # Table 6.3: ONT Troubleshooting Matrix (Chapter 6)
        data = [
            ("LOS (Red)", "Steady Glow", "Loss of Signal: Total physical fiber cut or severe macro-bend. Check outside plant with OTDR."),
            ("PON (Green)", "Blinking / Flashing", "Registering: Authentication in progress. If persistent, light power is weak or configuration profile is missing."),
            ("PON (Green)", "Steady Solid Glow", "Connected: Symmetrical operational link established successfully."),
            ("Power Levels", "Below -27 dBm", "Power Fault: Dirty connector faces or high-loss splice. Rework connection or clean with a OneClick cleaner.")
        ]
        headers = ["ONT Indicator", "Status State", "Meaning / Corrective Action"]
        caption = "Table 6.3: ONT Troubleshooting Matrix"
        
    elif table_type == "appliance_troubleshooting":
        # Chapter 10 Table: Troubleshooting & Maintenance Practice Matrix
        data = [
            ("Electric Iron", "No heating; indicator light off", "Open circuit in thermal fuse / Broken power cord", "Replaced thermal fuse; trimmed and re-terminated power cord."),
            ("Induction Cooker", "Dead display; won't power up", "Blown main fuse due to shorted IGBT", "Replaced the shorted IGBT and main glass fuse; checked gate driver circuit."),
            ("Mixer Grinder", "Sparks from bottom; sudden stop", "Worn carbon brushes / Overload switch tripped", "Replaced carbon brushes; cleared jar jam and reset the bottom OLP switch."),
            ("Water Heater", "Water not heating; indicator on", "Scale buildup or open-circuit in heating element", "Drained tank, removed limescale, and replaced the immersion heating element."),
            ("Ceiling Fan", "Humming sound; spins very slowly", "Degraded or leaked capacitor", "Desoldered/disconnected old capacitor, verified value, and installed new 2.5µF unit."),
            ("Wall Fan", "Fan rotates but does not oscillate / Overheating", "Stripped linkage gears / Dry rotor shaft", "Replaced the plastic gear wheel assembly; cleaned and lubricated the rotor shaft with machine oil."),
            ("Pump Motor", "Motor hums but does not rotate", "Seized bearings / Defective starter capacitor", "Freed impeller manually; replaced bearings and capacitor.")
        ]
        headers = ["Appliance", "Common Symptom", "Root Cause Detected", "Corrective Action Performed"]
        caption = "Table 10.1: Appliance Troubleshooting & Maintenance Practice Matrix"
        
    elif table_type == "activity_schedule":
        # Day-wise Internship Activity Log
        data = [
            ("Day 1", "21.05.2026", "Optical Fiber Theory Class (Introduction to OFC)"),
            ("Day 2", "22.05.2026", "Splicing Practical Section"),
            ("Day 3", "23.05.2026", "Repairing of the Ceiling Fan"),
            ("Day 4", "25.05.2026", "• Theory Class of Fiber To The Home (FTTH)\n• Repairing of Mixer Grinder\n• Field Visit of Splicing"),
            ("Day 5", "26.05.2026", "Field Visit at Manjeri Telephone Exchange"),
            ("Day 6", "08.06.2026", "Power Calculation Theory Class"),
            ("Day 7", "09.06.2026", "• Repairing of Iron Box\n• Repairing of Wall Fan\n• Repairing of Induction Cooker")
        ]
        headers = ["Day", "Date", "Topics / Activities Covered"]
        caption = "Table 1: Day-wise Internship Activity Log"
        
    else:
        return
        
    # Create Table
    num_rows = len(data) + 1
    num_cols = len(headers)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Format Header Row
    hdr_row = table.rows[0]
    for col_idx, text in enumerate(headers):
        cell = hdr_row.cells[col_idx]
        cell.text = text
        set_cell_background(cell, primary_hex)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        run = p.runs[0]
        run.font.name = font_name
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Format Data Rows
    for row_idx, row_data in enumerate(data):
        row = table.rows[row_idx + 1]
        
        # Zebra striping background color
        bg_hex = "F3F4F6" if row_idx % 2 == 1 else "FFFFFF"
        
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.text = text
            set_cell_background(cell, bg_hex)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            # Align first column center, rest left (unless it's color coding)
            if col_idx == 0 or table_type == "color_coding" or table_type == "activity_schedule" and col_idx < 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            if p.runs:
                run = p.runs[0]
                run.font.name = font_name
                run.font.size = Pt(9.5)
                
    # Add caption below the table
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption)
    run_cap.font.name = font_name
    run_cap.font.size = Pt(9)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(100, 100, 100)

# Build a cover page based on style
def build_cover_page(doc, student_name, theme):
    font_name = theme["font"]
    primary_color = theme["primary_color"]
    secondary_color = theme["secondary_color"]
    primary_hex = "{:02X}{:02X}{:02X}".format(*primary_color)
    
    style = theme["cover_style"]
    
    # Margin settings for first section (cover)
    section = doc.sections[0]
    section.top_margin = Inches(theme["margin_inches"])
    section.bottom_margin = Inches(theme["margin_inches"])
    section.left_margin = Inches(theme["margin_inches"])
    section.right_margin = Inches(theme["margin_inches"])
    
    if style == "minimalist":
        # Large title, colored horizontal bar, details centered below, logo at bottom
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(20)
            
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(10)
        run_title = p_title.add_run("INTERNSHIP REPORT\nON\nFIBER OPTIC COMMUNICATIONS, FTTH NETWORK ARCHITECTURE\nAND\nELECTRICAL HOME APPLIANCES MAINTENANCE")
        run_title.font.name = font_name
        run_title.font.size = Pt(16)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(*primary_color)
        
        # Color line
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, primary_hex)
        cell.width = Inches(4.5)
        tr = table.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), '40')
        trPr.append(trHeight)
        
        p_space2 = doc.add_paragraph()
        p_space2.paragraph_format.space_before = Pt(20)
            
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.line_spacing = 1.2
        p_sub.paragraph_format.space_after = Pt(20)
        
        r1 = p_sub.add_run("Submitted by:\n")
        r1.font.bold = True
        r2 = p_sub.add_run(f"{student_name}\n")
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(*secondary_color)
        r2.font.size = Pt(12)
        r3 = p_sub.add_run("In partial fulfillment of the requirements for the award of the Degree of\n")
        r4 = p_sub.add_run("Bachelor of Science in Physics\n\n")
        r4.font.bold = True
        r5 = p_sub.add_run("Under the Faculty Guidance of:\n")
        r5.font.bold = True
        r6 = p_sub.add_run("Dr. Mohemmed Shanid N A\nDepartment of Physics\nSullamussalam Science College, Areekode\n")
        
        for r in [r1, r2, r3, r4, r5, r6]:
            r.font.name = font_name
            if r != r2:
                r.font.size = Pt(10)
                
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(10)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.0))
            
        p_dept = doc.add_paragraph()
        p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_dept = p_dept.add_run("DEPARTMENT OF PHYSICS\nSULLAMUSSALAM SCIENCE COLLEGE, AREEKODE\nJune 2026")
        r_dept.font.name = font_name
        r_dept.font.size = Pt(9.5)
        r_dept.font.bold = True
        r_dept.font.color.rgb = RGBColor(*primary_color)
        
    elif style == "classic":
        # Logo at top, Title centered, details centered
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(10)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.1))
            
        p_college = doc.add_paragraph()
        p_college.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_college.paragraph_format.space_after = Pt(20)
        r_col = p_college.add_run("SULLAMUSSALAM SCIENCE COLLEGE, AREEKODE")
        r_col.font.name = font_name
        r_col.font.size = Pt(11.5)
        r_col.font.bold = True
        r_col.font.color.rgb = RGBColor(*primary_color)
        
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.line_spacing = 1.25
        p_title.paragraph_format.space_after = Pt(20)
        r_title = p_title.add_run("INTERNSHIP REPORT ON\nFIBER OPTIC COMMUNICATIONS, FTTH NETWORK ARCHITECTURE\nAND ELECTRICAL HOME APPLIANCES MAINTENANCE")
        r_title.font.name = font_name
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.line_spacing = 1.2
        p_sub.paragraph_format.space_after = Pt(20)
        
        r1 = p_sub.add_run("Candidate: ")
        r2 = p_sub.add_run(f"{student_name}\n")
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(*primary_color)
        r3 = p_sub.add_run("Reg No: BS23PHY07\n\n")
        r4 = p_sub.add_run("Submitted in partial fulfillment of the BSc Physics Degree\n")
        r5 = p_sub.add_run("Under the Academic Supervision of:\n")
        r6 = p_sub.add_run("Dr. Mohemmed Shanid N A\nHead of the Department of Physics\n")
        r6.font.bold = True
        
        for r in [r1, r2, r3, r4, r5, r6]:
            r.font.name = font_name
            r.font.size = Pt(10.5)
            
        p_assoc = doc.add_paragraph()
        p_assoc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_assoc = p_assoc.add_run("In Industrial Collaboration with:\nBfone Fibernet, Edavannappara\n& BSNL Exchange, Manjeri")
        r_assoc.font.name = font_name
        r_assoc.font.size = Pt(10)
        r_assoc.font.bold = True
        r_assoc.font.color.rgb = RGBColor(*secondary_color)
        
    elif style == "left_accent":
        # Two-column layout on cover: Left column is a thick primary color bar, Right column has content
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        remove_cell_borders(table.cell(0, 0))
        remove_cell_borders(table.cell(0, 1))
        
        # Left bar width
        cell_bar = table.cell(0, 0)
        cell_bar.width = Inches(0.3)
        set_cell_background(cell_bar, primary_hex)
        
        # Right text width
        cell_text = table.cell(0, 1)
        cell_text.width = Inches(5.5)
        
        # Add paragraphs to right cell
        p1 = cell_text.paragraphs[0]
        p1.paragraph_format.left_indent = Inches(0.25)
        p1.paragraph_format.space_before = Pt(10)
        p1.paragraph_format.space_after = Pt(10)
        run_rep = p1.add_run("INTERNSHIP REPORT")
        run_rep.font.name = font_name
        run_rep.font.size = Pt(13)
        run_rep.font.bold = True
        run_rep.font.color.rgb = RGBColor(*secondary_color)
        
        p2 = cell_text.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.25)
        p2.paragraph_format.space_after = Pt(30)
        p2.paragraph_format.line_spacing = 1.2
        run_t = p2.add_run("FIBER OPTIC COMMUNICATIONS,\nFTTH NETWORK ARCHITECTURE\nAND HOME APPLIANCES")
        run_t.font.name = font_name
        run_t.font.size = Pt(17)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(*primary_color)
        
        p3 = cell_text.add_paragraph()
        p3.paragraph_format.left_indent = Inches(0.25)
        p3.paragraph_format.space_after = Pt(35)
        p3.paragraph_format.line_spacing = 1.25
        
        runs = [
            p3.add_run("SUBMITTED BY:\n"),
            p3.add_run(f"{student_name}\n\n"),
            p3.add_run("SUPERVISED BY:\n"),
            p3.add_run("Dr. Mohemmed Shanid N A\n\n"),
            p3.add_run("INSTITUTION:\n"),
            p3.add_run("Department of Physics\nSullamussalam Science College, Areekode\n")
        ]
        runs[1].font.bold = True
        runs[1].font.color.rgb = RGBColor(*primary_color)
        runs[3].font.bold = True
        runs[5].font.bold = True
        
        for run in runs:
            run.font.name = font_name
            run.font.size = Pt(10)
            
        p4 = cell_text.add_paragraph()
        p4.paragraph_format.left_indent = Inches(0.25)
        p4.paragraph_format.space_before = Pt(10)
        if os.path.exists(LOGO_PATH):
            p4.add_run().add_picture(LOGO_PATH, width=Inches(0.95))
            
    elif style == "double_border":
        # A single-cell table with double borders that fits exactly
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(5.8)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="double" w:sz="12" w:space="0" w:color="{primary_hex}"/>
                <w:left w:val="double" w:sz="12" w:space="0" w:color="{primary_hex}"/>
                <w:bottom w:val="double" w:sz="12" w:space="0" w:color="{primary_hex}"/>
                <w:right w:val="double" w:sz="12" w:space="0" w:color="{primary_hex}"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        # Add content inside the cell
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(20)
        
        run_rep = p.add_run("INTERNSHIP REPORT\n")
        run_rep.font.name = font_name
        run_rep.font.size = Pt(13)
        run_rep.font.bold = True
        run_rep.font.color.rgb = RGBColor(*secondary_color)
        
        run_t = p.add_run("\nFIBER OPTIC COMMUNICATIONS,\nFTTH NETWORK ARCHITECTURE\nAND ELECTRICAL HOME APPLIANCES\nMAINTENANCE\n")
        run_t.font.name = font_name
        run_t.font.size = Pt(15)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(*primary_color)
        
        if os.path.exists(LOGO_PATH):
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_after = Pt(20)
            p_img.add_run().add_picture(LOGO_PATH, width=Inches(1.0))
            
        p_det = cell.add_paragraph()
        p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_det.paragraph_format.line_spacing = 1.2
        p_det.paragraph_format.space_after = Pt(20)
        
        runs = [
            p_det.add_run("Submitted by:\n"),
            p_det.add_run(f"{student_name}\n\n"),
            p_det.add_run("BSc Physics Student\n\n"),
            p_det.add_run("Under the guidance of:\n"),
            p_det.add_run("Dr. Mohemmed Shanid N A\n"),
            p_det.add_run("Department of Physics\nSullamussalam Science College, Areekode")
        ]
        runs[1].font.bold = True
        runs[1].font.color.rgb = RGBColor(*primary_color)
        runs[1].font.size = Pt(11.5)
        runs[4].font.bold = True
        
        for run in runs:
            run.font.name = font_name
            if run != runs[1]:
                run.font.size = Pt(10)
                
    elif style == "split_title":
        # Shaded title block at top, details below
        table = doc.add_table(rows=2, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        remove_cell_borders(table.cell(0, 0))
        remove_cell_borders(table.cell(1, 0))
        
        cell_top = table.cell(0, 0)
        cell_top.width = Inches(5.8)
        set_cell_background(cell_top, primary_hex)
        
        p_top = cell_top.paragraphs[0]
        p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_top.paragraph_format.space_before = Pt(20)
        p_top.paragraph_format.space_after = Pt(20)
        
        r_title = p_top.add_run("INTERNSHIP REPORT\nON\nFIBER OPTIC COMMUNICATIONS, FTTH NETWORK ARCHITECTURE\nAND ELECTRICAL HOME APPLIANCES MAINTENANCE")
        r_title.font.name = font_name
        r_title.font.size = Pt(14)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(255, 255, 255)
        
        cell_bottom = table.cell(1, 0)
        cell_bottom.width = Inches(5.8)
        
        p_logo = cell_bottom.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(30)
        p_logo.paragraph_format.space_after = Pt(20)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.0))
            
        p_det = cell_bottom.add_paragraph()
        p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_det.paragraph_format.line_spacing = 1.2
        p_det.paragraph_format.space_after = Pt(20)
        
        runs = [
            p_det.add_run("Submitted by:\n"),
            p_det.add_run(f"{student_name}\n\n"),
            p_det.add_run("Department of Physics\nSullamussalam Science College, Areekode\n\n"),
            p_det.add_run("Under the Faculty Guidance of:\n"),
            p_det.add_run("Dr. Mohemmed Shanid N A\nDepartment of Physics")
        ]
        runs[1].font.bold = True
        runs[1].font.color.rgb = RGBColor(*primary_color)
        runs[4].font.bold = True
        
        for run in runs:
            run.font.name = font_name
            run.font.size = Pt(10)
            
    elif style == "elegant_center":
        # Centered layout with horizontal divider lines above and below title
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(20)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.0))
            
        # Top line
        table1 = doc.add_table(rows=1, cols=1)
        table1.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_background(table1.cell(0, 0), primary_hex)
        table1.cell(0, 0).width = Inches(5.0)
        tr1 = table1.rows[0]._tr
        trPr1 = tr1.get_or_add_trPr()
        trHeight1 = OxmlElement('w:trHeight')
        trHeight1.set(qn('w:val'), '15')
        trPr1.append(trHeight1)
        
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(10)
        p_title.paragraph_format.space_after = Pt(10)
        p_title.paragraph_format.line_spacing = 1.25
        r_title = p_title.add_run("INTERNSHIP REPORT\nFIBER OPTIC COMMUNICATIONS, FTTH NETWORK ARCHITECTURE\nAND HOME APPLIANCES MAINTENANCE")
        r_title.font.name = font_name
        r_title.font.size = Pt(14.5)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(*secondary_color)
        
        # Bottom line
        table2 = doc.add_table(rows=1, cols=1)
        table2.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_cell_background(table2.cell(0, 0), primary_hex)
        table2.cell(0, 0).width = Inches(5.0)
        tr2 = table2.rows[0]._tr
        trPr2 = tr2.get_or_add_trPr()
        trHeight2 = OxmlElement('w:trHeight')
        trHeight2.set(qn('w:val'), '15')
        trPr2.append(trHeight2)
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(20)
        
        p_det = doc.add_paragraph()
        p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_det.paragraph_format.line_spacing = 1.2
        p_det.paragraph_format.space_after = Pt(20)
        
        runs = [
            p_det.add_run("Submitted by:\n"),
            p_det.add_run(f"{student_name}\n\n"),
            p_det.add_run("Under the Faculty Guidance of:\n"),
            p_det.add_run("Dr. Mohemmed Shanid N A\n\n"),
            p_det.add_run("DEPARTMENT OF PHYSICS\nSULLAMUSSALAM SCIENCE COLLEGE, AREEKODE")
        ]
        runs[1].font.bold = True
        runs[1].font.color.rgb = RGBColor(*primary_color)
        runs[3].font.bold = True
        runs[4].font.bold = True
        
        for run in runs:
            run.font.name = font_name
            run.font.size = Pt(10)
            
    elif style == "geometric":
        # Geometric layout: Title inside box, details inside borderless table below
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.width = Inches(5.5)
        set_cell_background(cell, primary_hex)
        
        p_box = cell.paragraphs[0]
        p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_box.paragraph_format.space_before = Pt(15)
        p_box.paragraph_format.space_after = Pt(15)
        
        r_t1 = p_box.add_run("INTERNSHIP REPORT\n")
        r_t1.font.name = font_name
        r_t1.font.size = Pt(13)
        r_t1.font.bold = True
        r_t1.font.color.rgb = RGBColor(*secondary_color)
        
        r_t2 = p_box.add_run("OFC, FTTH AND ELECTRICAL APPLIANCES")
        r_t2.font.name = font_name
        r_t2.font.size = Pt(15)
        r_t2.font.bold = True
        r_t2.font.color.rgb = RGBColor(255, 255, 255)
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(25)
        
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(25)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.0))
            
        p_det = doc.add_paragraph()
        p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_det.paragraph_format.line_spacing = 1.25
        p_det.paragraph_format.space_after = Pt(20)
        
        runs = [
            p_det.add_run("Candidate Name: "),
            p_det.add_run(f"{student_name}\n"),
            p_det.add_run("Department: Physics\n"),
            p_det.add_run("Academic Supervisor: Dr. Mohemmed Shanid N A\n"),
            p_det.add_run("Institution: Sullamussalam Science College, Areekode\n")
        ]
        runs[1].font.bold = True
        runs[1].font.color.rgb = RGBColor(*primary_color)
        
        for run in runs:
            run.font.name = font_name
            run.font.size = Pt(10)
            
    else:
        # Default standard centered layout with logo in center, colored lines
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(10)
            
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_after = Pt(15)
        r_title = p_title.add_run("INTERNSHIP REPORT\nON\nFIBER OPTIC COMMUNICATIONS, FTTH NETWORK ARCHITECTURE\nAND ELECTRICAL HOME APPLIANCES MAINTENANCE")
        r_title.font.name = font_name
        r_title.font.size = Pt(15)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(*primary_color)
        
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(15)
        if os.path.exists(LOGO_PATH):
            p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.1))
            
        p_details = doc.add_paragraph()
        p_details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_details.paragraph_format.line_spacing = 1.25
        p_details.paragraph_format.space_after = Pt(15)
        
        r1 = p_details.add_run("Submitted by:\n")
        r2 = p_details.add_run(f"{student_name}\n")
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(*secondary_color)
        r2.font.size = Pt(11.5)
        r3 = p_details.add_run("Bachelor of Science in Physics\n\n")
        r3.font.bold = True
        r4 = p_details.add_run("Faculty Guide:\n")
        r5 = p_details.add_run("Dr. Mohemmed Shanid N A\nDepartment of Physics\nSullamussalam Science College, Areekode\n")
        r5.font.bold = True
        
        for r in [r1, r2, r3, r4, r5]:
            r.font.name = font_name
            if r != r2:
                r.font.size = Pt(10)

# Heading levels with styles
def add_heading_styled(doc, text, level, theme):
    font_name = theme["font"]
    primary_color = theme["primary_color"]
    secondary_color = theme["secondary_color"]
    
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 0:
        # Part title page or huge section
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(100)
        p.paragraph_format.space_after = Pt(20)
        run = p.add_run(text)
        run.font.name = font_name
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*primary_color)
    elif level == 1:
        # Chapter title
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(12)
        
        # Accent left border or block can be simulated in Word, but simple coloring is clean.
        run = p.add_run(text.upper())
        run.font.name = font_name
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*primary_color)
        
        # Add a subtle line below the chapter heading
        # table = doc.add_table(rows=1, cols=1)
        # table.alignment = WD_TABLE_ALIGNMENT.LEFT
    elif level == 2:
        # Section
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = font_name
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*primary_color)
    elif level == 3:
        # Subsection
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = font_name
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(*secondary_color)
        
    return p

# Write Certificate Page
def write_certificate_page(doc, student_name, theme):
    font_name = theme["font"]
    primary_color = theme["primary_color"]
    
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert.paragraph_format.space_before = Pt(20)
    p_cert.paragraph_format.space_after = Pt(30)
    run_cert = p_cert.add_run("CERTIFICATE")
    run_cert.font.name = font_name
    run_cert.font.size = Pt(16)
    run_cert.font.bold = True
    run_cert.font.color.rgb = RGBColor(*primary_color)
    
    cert_text = (
        f"This is to certify that the internship report entitled ”Fiber Optic Communications, FTTH Network "
        f"Architecture and Electrical Home Appliances Maintenance” is a bona fide record of the industrial training "
        f"carried out by {student_name} under my academic supervision and guidance, in corporate collaboration with "
        f"Bfone Fibernet, Edavannappara, in partial fulfillment of the requirements for the award of the Degree of "
        f"Bachelor of Science in Physics."
    )
    
    p_body = doc.add_paragraph()
    p_body.paragraph_format.line_spacing = 1.3
    p_body.paragraph_format.space_after = Pt(60)
    p_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run_body = p_body.add_run(cert_text)
    run_body.font.name = font_name
    run_body.font.size = Pt(11.5)
    
    # Signature Lines Table
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    remove_cell_borders(table.cell(0, 0))
    remove_cell_borders(table.cell(0, 1))
    remove_cell_borders(table.cell(1, 0))
    remove_cell_borders(table.cell(1, 1))
    
    # Row 0: Signatures
    c0 = table.cell(0, 0)
    c0.width = Inches(3.0)
    p_sig1 = c0.paragraphs[0]
    r_sig1 = p_sig1.add_run("Dr. Mohemmed Shanid N A\nHead of the Department\nDepartment of Physics")
    r_sig1.font.name = font_name
    r_sig1.font.size = Pt(10.5)
    r_sig1.font.bold = True
    p_sig1.paragraph_format.line_spacing = 1.15
    
    c1 = table.cell(0, 1)
    c1.width = Inches(3.0)
    p_sig2 = c1.paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig2 = p_sig2.add_run("External Examiner")
    r_sig2.font.name = font_name
    r_sig2.font.size = Pt(10.5)
    r_sig2.font.bold = True
    
    # Row 1
    c2 = table.cell(1, 0)
    p_sig3 = c2.paragraphs[0]
    p_sig3.paragraph_format.space_before = Pt(40)
    r_sig3 = p_sig3.add_run("Internal Examiner")
    r_sig3.font.name = font_name
    r_sig3.font.size = Pt(10.5)
    r_sig3.font.bold = True

# Write Table of Contents Page
def write_toc_page(doc, theme, page_numbers=None):
    font_name = theme["font"]
    primary_color = theme["primary_color"]
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(20)
    run_t = p_title.add_run("CONTENTS")
    run_t.font.name = font_name
    run_t.font.size = Pt(16)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(*primary_color)
    
    # List of entries in TOC
    toc_entries = [
        ("CERTIFICATE", "CERTIFICATE", 0),
        ("INTERNSHIP ACTIVITY SCHEDULE", "INTERNSHIP ACTIVITY SCHEDULE", 0),
        ("Part I: Fiber Optic Communications and FTTH Network Architecture", "Part I", 0),
        ("Part Introduction: Fiber Optic Infrastructure", "Part Introduction: Fiber Optic Infrastructure", 1),
        ("1 Introduction to Optical Fiber Communication", "Chapter 1", 1),
        ("1.1 Overview", "1.1", 2),
        ("1.2 The Basic Fiber Optic Link", "1.2", 2),
        ("1.3 Advantages of Fiber Optics", "1.3", 2),
        ("2 Principle and Cable Structure", "Chapter 2", 1),
        ("2.1 Principle of Operation: Total Internal Reflection", "2.1", 2),
        ("2.2 Cable Anatomy", "2.2", 2),
        ("2.3 Cable Types", "2.3", 2),
        ("3 Link Budgeting, Splicing, and Mathematical Calculations", "Chapter 3", 1),
        ("3.1 OFC Splicing", "3.1", 2),
        ("3.2 Logarithmic Conversion Mechanics", "3.2", 2),
        ("3.3 Splitter Loss Calculations", "3.3", 2),
        ("3.4 Standard Loss Values for Link Performance", "3.4", 2),
        ("3.5 Mathematical Verification: Power Link Budget Case Study", "3.5 ", 2),
        ("3.5.1 Case Study Parameter Blueprint", "3.5.1", 3),
        ("3.5.2 Link Status Analysis", "3.5.2", 3),
        ("4 Optical Fiber Splicing: Detailed Analysis", "Chapter 4", 1),
        ("4.1 What is Splicing?", "4.1", 2),
        ("4.2 Types of Splicing", "4.2", 2),
        ("4.2.1 Fusion Splicing", "4.2.1", 3),
        ("4.2.2 Mechanical Splicing", "4.2.2", 3),
        ("4.3 Importance and Significance", "4.3", 2),
        ("4.4 Step-by-Step Fusion Splicing Process", "4.4", 2),
        ("4.5 Advantages", "4.5", 2),
        ("4.6 Future Scope", "4.6", 2),
        ("4.7 Conclusion", "4.7", 2),
        ("5 Introduction to FTTH & Passive Optical Networks (PON)", "Chapter 5", 1),
        ("5.1 Overview of FTTH", "5.1", 2),
        ("5.2 The FTTH Network Architecture Layers", "5.2", 2),
        ("5.2.1 A. The Feeder Segment", "5.2.1", 3),
        ("5.2.2 B. The Distribution Segment", "5.2.2", 3),
        ("5.2.3 C. The Drop Segment", "5.2.3", 3),
        ("5.3 Core Hardware Components", "5.3", 2),
        ("5.3.1 Active Components (Power Required)", "5.3.1", 3),
        ("5.3.2 Passive Components (No Power Required)", "5.3.2", 3),
        ("6 FTTH Deployment, Coding and Testing", "Chapter 6", 1),
        ("6.1 Fiber Color Coding Standard (TIA-598-C)", "6.1", 2),
        ("6.2 Physical Deployment & Site Survey", "6.2", 2),
        ("6.2.1 Site Survey & Planning", "6.2.1", 3),
        ("6.2.2 Outdoor Routing Operations", "6.2.2", 3),
        ("6.2.3 Indoor Termination", "6.2.3", 3),
        ("6.3 Splicing Methods Comparison", "6.3", 2),
        ("6.4 Link Budget & Testing Benchmarks", "6.4", 2),
        ("6.4.1 Wavelength Multiplexing", "6.4.1", 3),
        ("6.4.2 Testing Parameters", "6.4.2", 3),
        ("6.5 Troubleshooting & Indicators", "6.5", 2),
        ("7 Field Visit: BSNL Exchange Office, Manjeri", "Chapter 7", 1),
        ("7.1 Introduction to the Field Visit", "7.1", 2),
        ("7.2 Observations and Core Infrastructure", "7.2", 2),
        ("7.3 Conclusion of the Visit", "7.3", 2),
        ("Part Conclusion: Fiber Optic Infrastructure", "Part Conclusion: Fiber Optic Infrastructure", 1),
        ("Part II: Electrical Home Appliances Maintenance and Troubleshooting", "Part II", 0),
        ("Part Introduction: Domestic Electrical Engineering", "Part Introduction: Domestic Electrical Engineering", 1),
        ("8 Fundamentals of Electrical Systems & Safety", "Chapter 8", 1),
        ("8.1 Introduction", "8.1", 2),
        ("8.2 Fundamentals of Electrical Systems & Safety Overview", "8.2", 2),
        ("8.2.1 Circuit Behaviors & Laws", "8.2.1", 3),
        ("8.2.2 Earthing and Protection Systems", "8.2.2", 3),
        ("9 Technical Breakdown & Diagnostics of Core Home Appliances", "Chapter 9", 1),
        ("9.1 Electric Iron", "9.1", 2),
        ("9.2 Induction Cooker", "9.2", 2),
        ("9.3 Mixer Grinder", "9.3 ", 2),
        ("9.3.1 Ceiling Fan", "9.3.1", 3),
        ("9.3.2 Wall Fan", "9.3.2", 3),
        ("Chapter 10: Troubleshooting & Maintenance Practice Matrix", "Chapter 10", 1),
        ("10.1 Internship Outcomes & Skills Acquired", "10.1", 2),
        ("11 Applied Practical Calculations", "Chapter 11", 1),
        ("11.1 Domestic Electricity Consumption Auditing", "11.1", 2),
        ("11.2 Charge Calculation in Audio Systems", "11.2", 2),
        ("Part Conclusion: Domestic Electrical Engineering", "Part Conclusion: Domestic Electrical Engineering", 1)
    ]
    
    # Calculate right-align tab stop position based on A4 width (8.27) and margins
    margin_in = theme["margin_inches"]
    tab_stop_pos = 8.27 - (2 * margin_in) - 0.05
    
    for label, search_key, level in toc_entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(level * 0.25)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(tab_stop_pos), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        # Style part titles differently in TOC
        is_part = label.startswith("Part I:") or label.startswith("Part II:")
        
        run_lbl = p.add_run(label)
        run_lbl.font.name = font_name
        run_lbl.font.size = Pt(10 if is_part else 9.5)
        run_lbl.font.bold = is_part
        if is_part:
            run_lbl.font.color.rgb = RGBColor(*primary_color)
            
        p.add_run("\t")
        
        page_val = "XX"
        if page_numbers and search_key in page_numbers:
            page_val = str(page_numbers[search_key])
            
        run_pg = p.add_run(page_val)
        run_pg.font.name = font_name
        run_pg.font.size = Pt(9.5)
        run_pg.font.bold = True
        if is_part:
            run_pg.font.color.rgb = RGBColor(*primary_color)

# Add custom footers with page numbers
def add_footer_page_number(doc, student_name, theme):
    def add_p_border_top(p, color_hex="CCCCCC", size=4):
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        top = OxmlElement('w:top')
        top.set(qn('w:val'), 'single')
        top.set(qn('w:sz'), str(size))  # 4 = 0.5 pt
        top.set(qn('w:space'), '6')     # spacing
        top.set(qn('w:color'), color_hex)
        pBdr.append(top)
        pPr.append(pBdr)

    # Word supports footer page numbers via fields. We can add a simple page number field XML.
    for i, section in enumerate(doc.sections):
        if i == 0:
            # First section is Cover Page, skip it
            continue
        
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.text = "" # clear default
        
        # Add border
        add_p_border_top(p, color_hex="D1D5DB", size=4)
        
        font_name = theme["font"]
        
        run_text = p.add_run(f"{student_name}  |  BSc Physics Internship Report  |  Page ")
        run_text.font.name = font_name
        run_text.font.size = Pt(8.5)
        run_text.font.color.rgb = RGBColor(120, 120, 120)
        
        run_num = p.add_run()
        run_num.font.name = font_name
        run_num.font.size = Pt(8.5)
        run_num.font.bold = True
        run_num.font.color.rgb = RGBColor(120, 120, 120)
        
        # Insert Page field XML inside run_num._r
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        
        run_num._r.append(fldChar1)
        run_num._r.append(instrText)
        run_num._r.append(fldChar2)
        run_num._r.append(fldChar3)


# Build Document
def build_report_docx(student_name, theme, student_index, page_numbers=None):
    doc = Document()
    font_name = theme["font"]
    primary_color = theme["primary_color"]
    secondary_color = theme["secondary_color"]
    
    # 1. Cover Page
    build_cover_page(doc, student_name, theme)
    
    # 2. Certificate Section (New Section so we can set headers/footers)
    doc.add_section()
    doc.sections[-1].header.is_linked_to_previous = False
    doc.sections[-1].footer.is_linked_to_previous = False
    
    # Apply margins to section 2
    doc.sections[-1].top_margin = Inches(theme["margin_inches"])
    doc.sections[-1].bottom_margin = Inches(theme["margin_inches"])
    doc.sections[-1].left_margin = Inches(theme["margin_inches"])
    doc.sections[-1].right_margin = Inches(theme["margin_inches"])
    
    write_certificate_page(doc, student_name, theme)
    
    # 3. Table of Contents Section
    doc.add_page_break()
    write_toc_page(doc, theme, page_numbers)
    
    # 4. Activity Schedule Section
    doc.add_page_break()
    add_heading_styled(doc, "INTERNSHIP ACTIVITY SCHEDULE", 1, theme)
    doc.add_paragraph()
    build_styled_table(doc, "activity_schedule", theme, student_index)
    
    # 5. Part I Section
    doc.add_page_break()
    add_heading_styled(doc, "Part I\nFiber Optic Communications and\nFTTH Network Architecture", 0, theme)
    
    # Part I Intro
    doc.add_page_break()
    add_heading_styled(doc, "Part Introduction: Fiber Optic Infrastructure", 1, theme)
    add_p(doc, REPORT_CONTENT["part1_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Chapter 1
    add_heading_styled(doc, "Chapter 1\nIntroduction to Optical Fiber Communication", 1, theme)
    add_heading_styled(doc, "1.1 Overview", 2, theme)
    add_p(doc, REPORT_CONTENT["ch1_overview"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "1.2 The Basic Fiber Optic Link", 2, theme)
    add_p(doc, REPORT_CONTENT["ch1_link_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for item in REPORT_CONTENT["ch1_link_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "1.3 Advantages of Fiber Optics", 2, theme)
    for item in REPORT_CONTENT["ch1_advantages"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    # Chapter 2
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 2\nPrinciple and Cable Structure", 1, theme)
    add_heading_styled(doc, "2.1 Principle of Operation: Total Internal Reflection", 2, theme)
    add_p(doc, REPORT_CONTENT["ch2_tir"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for item in REPORT_CONTENT["ch2_tir_conditions"]:
        add_list_item(doc, item, student_index, font_name, 10.5, bullet_char="• ")
        
    # Formula
    add_p(doc, REPORT_CONTENT["ch2_tir_formula"], student_index, font_name, 11)
    p_eq = doc.add_paragraph()
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq.paragraph_format.space_before = Pt(6)
    p_eq.paragraph_format.space_after = Pt(12)
    run_eq = p_eq.add_run("θc = sin^-1( n2 / n1 )     (2.1)")
    run_eq.font.name = font_name
    run_eq.font.size = Pt(11.5)
    run_eq.font.bold = True
    
    add_heading_styled(doc, "2.2 Cable Anatomy", 2, theme)
    add_p(doc, REPORT_CONTENT["ch2_anatomy_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for item in REPORT_CONTENT["ch2_anatomy_layers"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "2.3 Cable Types", 2, theme)
    for item in REPORT_CONTENT["ch2_cable_types"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    # Chapter 3
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 3\nLink Budgeting, Splicing, and Mathematical Calculations", 1, theme)
    add_heading_styled(doc, "3.1 OFC Splicing", 2, theme)
    add_p(doc, REPORT_CONTENT["ch3_splicing_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for item in REPORT_CONTENT["ch3_splicing_types"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "3.2 Logarithmic Conversion Mechanics", 2, theme)
    add_p(doc, REPORT_CONTENT["ch3_log_conv"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_eq2 = doc.add_paragraph()
    p_eq2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq2.paragraph_format.space_before = Pt(6)
    p_eq2.paragraph_format.space_after = Pt(12)
    run_eq2 = p_eq2.add_run("P_dBm = 10 log10( P_mW / 1 mW )     (3.1)")
    run_eq2.font.name = font_name
    run_eq2.font.size = Pt(11.5)
    run_eq2.font.bold = True
    
    add_heading_styled(doc, "3.3 Splitter Loss Calculations", 2, theme)
    add_p(doc, REPORT_CONTENT["ch3_splitter_loss"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_eq3 = doc.add_paragraph()
    p_eq3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq3.paragraph_format.space_before = Pt(6)
    p_eq3.paragraph_format.space_after = Pt(12)
    run_eq3 = p_eq3.add_run("Loss_Ideal = 10 log10( N )     (3.2)")
    run_eq3.font.name = font_name
    run_eq3.font.size = Pt(11.5)
    run_eq3.font.bold = True
    
    add_heading_styled(doc, "3.4 Standard Loss Values for Link Performance", 2, theme)
    add_p(doc, REPORT_CONTENT["ch3_standard_loss_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    build_styled_table(doc, "loss_parameters", theme, student_index)
    
    add_heading_styled(doc, "3.5 Mathematical Verification: Power Link Budget Case Study", 2, theme)
    p_eq4 = doc.add_paragraph()
    p_eq4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq4.paragraph_format.space_before = Pt(4)
    p_eq4.paragraph_format.space_after = Pt(8)
    run_eq4 = p_eq4.add_run("Pr = Pt - Loss_total     (3.3)")
    run_eq4.font.name = font_name
    run_eq4.font.size = Pt(11.5)
    run_eq4.font.bold = True
    
    add_heading_styled(doc, "3.5.1 Case Study Parameter Blueprint", 3, theme)
    add_p(doc, REPORT_CONTENT["ch3_case_study_blueprint"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for item in REPORT_CONTENT["ch3_case_study_list"]:
        add_list_item(doc, item, student_index, font_name, 10.5, bullet_char="• ")
        
    add_p(doc, REPORT_CONTENT["ch3_case_study_sum"], student_index, font_name, 11)
    p_calc = doc.add_paragraph()
    p_calc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_calc.paragraph_format.left_indent = Inches(0.5)
    p_calc.paragraph_format.space_before = Pt(4)
    p_calc.paragraph_format.space_after = Pt(8)
    run_calc = p_calc.add_run("Total Attenuation Loss = 2.8 + 13.0 + 10.0 + 0.4 + 0.7\n                        = 26.9 dB")
    run_calc.font.name = font_name
    run_calc.font.size = Pt(11)
    run_calc.font.bold = True
    
    add_heading_styled(doc, "3.5.2 Link Status Analysis", 3, theme)
    add_p(doc, REPORT_CONTENT["ch3_link_status"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_eq5 = doc.add_paragraph()
    p_eq5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq5.paragraph_format.space_before = Pt(4)
    p_eq5.paragraph_format.space_after = Pt(8)
    run_eq5 = p_eq5.add_run("Receiver Sensitivity Threshold ≤ -28 dBm     (3.4)")
    run_eq5.font.name = font_name
    run_eq5.font.size = Pt(11)
    run_eq5.font.bold = True
    
    add_p(doc, REPORT_CONTENT["ch3_link_status_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Chapter 4
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 4\nOptical Fiber Splicing: Detailed Analysis", 1, theme)
    add_heading_styled(doc, "4.1 What is Splicing?", 2, theme)
    add_p(doc, REPORT_CONTENT["ch4_splicing_def"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "4.2 Types of Splicing", 2, theme)
    add_p(doc, REPORT_CONTENT["ch4_types_intro"], student_index, font_name, 11)
    
    add_heading_styled(doc, "4.2.1 Fusion Splicing", 3, theme)
    p_fus = doc.add_paragraph()
    p_fus.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_fus = p_fus.add_run(parse_spintax(REPORT_CONTENT["ch4_fusion_details"], student_index))
    r_fus.font.name = font_name
    r_fus.font.size = Pt(11)
    
    add_heading_styled(doc, "4.2.2 Mechanical Splicing", 3, theme)
    p_mec = doc.add_paragraph()
    p_mec.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_mec = p_mec.add_run(parse_spintax(REPORT_CONTENT["ch4_mech_details"], student_index))
    r_mec.font.name = font_name
    r_mec.font.size = Pt(11)
    
    add_heading_styled(doc, "4.3 Importance and Significance", 2, theme)
    for item in REPORT_CONTENT["ch4_importance"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "4.4 Step-by-Step Fusion Splicing Process", 2, theme)
    for i, item in enumerate(REPORT_CONTENT["ch4_steps"]):
        add_list_item(doc, item, student_index, font_name, 10.5, bullet_char=f"{i+1}. ", left_indent=0.3)
        
    add_heading_styled(doc, "4.5 Advantages", 2, theme)
    for item in REPORT_CONTENT["ch4_advantages"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "4.6 Future Scope", 2, theme)
    for item in REPORT_CONTENT["ch4_future"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "4.7 Conclusion", 2, theme)
    add_p(doc, REPORT_CONTENT["ch4_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Chapter 5
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 5\nIntroduction to FTTH & Passive Optical Networks (PON)", 1, theme)
    add_heading_styled(doc, "5.1 Overview of FTTH", 2, theme)
    add_p(doc, REPORT_CONTENT["ch5_ftth_overview"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "5.2 The FTTH Network Architecture Layers", 2, theme)
    add_p(doc, REPORT_CONTENT["ch5_layers_intro"], student_index, font_name, 11)
    for item in REPORT_CONTENT["ch5_layers"]:
        p_lay = doc.add_paragraph()
        p_lay.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_lay.paragraph_format.left_indent = Inches(0.25)
        p_lay.paragraph_format.space_after = Pt(6)
        r_lay = p_lay.add_run(parse_spintax(item, student_index))
        r_lay.font.name = font_name
        r_lay.font.size = Pt(11)
        
    add_heading_styled(doc, "5.3 Core Hardware Components", 2, theme)
    add_p(doc, REPORT_CONTENT["ch5_hardware_intro"], student_index, font_name, 11)
    
    add_heading_styled(doc, "5.3.1 Active Components (Power Required)", 3, theme)
    for item in REPORT_CONTENT["ch5_active_hardware"]:
        p_act = doc.add_paragraph()
        p_act.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_act.paragraph_format.left_indent = Inches(0.25)
        p_act.paragraph_format.space_after = Pt(6)
        r_act = p_act.add_run(parse_spintax(item, student_index))
        r_act.font.name = font_name
        r_act.font.size = Pt(11)
        
    add_heading_styled(doc, "5.3.2 Passive Components (No Power Required)", 3, theme)
    for item in REPORT_CONTENT["ch5_passive_hardware"]:
        p_pas = doc.add_paragraph()
        p_pas.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_pas.paragraph_format.left_indent = Inches(0.25)
        p_pas.paragraph_format.space_after = Pt(6)
        r_pas = p_pas.add_run(parse_spintax(item, student_index))
        r_pas.font.name = font_name
        r_pas.font.size = Pt(11)
        
    # Chapter 6
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 6\nFTTH Deployment, Coding and Testing", 1, theme)
    add_heading_styled(doc, "6.1 Fiber Color Coding Standard (TIA-598-C)", 2, theme)
    add_p(doc, REPORT_CONTENT["ch6_coding_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    build_styled_table(doc, "color_coding", theme, student_index)
    
    add_heading_styled(doc, "6.2 Physical Deployment & Site Survey", 2, theme)
    add_heading_styled(doc, "6.2.1 Site Survey & Planning", 3, theme)
    add_p(doc, REPORT_CONTENT["ch6_deployment_intro"], student_index, font_name, 11)
    for item in REPORT_CONTENT["ch6_deployment_steps"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "6.2.2 Outdoor Routing Operations", 3, theme)
    add_p(doc, REPORT_CONTENT["ch6_outdoor_routing"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "6.2.3 Indoor Termination", 3, theme)
    add_p(doc, REPORT_CONTENT["ch6_indoor_termination"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "6.3 Splicing Methods Comparison", 2, theme)
    build_styled_table(doc, "splicing_comparison", theme, student_index)
    
    add_heading_styled(doc, "6.4 Link Budget & Testing Benchmarks", 2, theme)
    add_p(doc, REPORT_CONTENT["ch6_testing_intro"], student_index, font_name, 11)
    
    add_heading_styled(doc, "6.4.1 Wavelength Multiplexing", 3, theme)
    for item in REPORT_CONTENT["ch6_wavelengths"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "6.4.2 Testing Parameters", 3, theme)
    for item in REPORT_CONTENT["ch6_testing_params"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "6.5 Troubleshooting & Indicators", 2, theme)
    add_p(doc, REPORT_CONTENT["ch6_troubleshoot_intro"], student_index, font_name, 11)
    build_styled_table(doc, "troubleshooting_matrix", theme, student_index)
    
    # Chapter 7
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 7\nField Visit: BSNL Exchange Office, Manjeri", 1, theme)
    add_heading_styled(doc, "7.1 Introduction to the Field Visit", 2, theme)
    add_p(doc, REPORT_CONTENT["ch7_visit_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "7.2 Observations and Core Infrastructure", 2, theme)
    add_p(doc, REPORT_CONTENT["ch7_observations"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "7.3 Conclusion of the Visit", 2, theme)
    add_p(doc, REPORT_CONTENT["ch7_visit_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Part 1 Conclusion
    add_heading_styled(doc, "Part Conclusion: Fiber Optic Infrastructure", 1, theme)
    add_p(doc, REPORT_CONTENT["part1_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Part II Section
    doc.add_page_break()
    add_heading_styled(doc, "Part II\nElectrical Home Appliances\nMaintenance and Troubleshooting", 0, theme)
    
    # Part II Intro
    doc.add_page_break()
    add_heading_styled(doc, "Part Introduction: Domestic Electrical Engineering", 1, theme)
    add_p(doc, REPORT_CONTENT["part2_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Chapter 8
    add_heading_styled(doc, "Chapter 8\nFundamentals of Electrical Systems & Safety", 1, theme)
    add_heading_styled(doc, "8.1 Introduction", 2, theme)
    add_p(doc, REPORT_CONTENT["ch8_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "8.2 Fundamentals of Electrical Systems & Safety Overview", 2, theme)
    add_p(doc, REPORT_CONTENT["ch8_overview"], student_index, font_name, 11)
    
    add_heading_styled(doc, "8.2.1 Circuit Behaviors & Laws", 3, theme)
    for item in REPORT_CONTENT["ch8_behaviors"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    # Formula
    p_eq6 = doc.add_paragraph()
    p_eq6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq6.paragraph_format.space_before = Pt(4)
    p_eq6.paragraph_format.space_after = Pt(8)
    run_eq6 = p_eq6.add_run("V = IR   and   P = VI = I^2R     (8.1)")
    run_eq6.font.name = font_name
    run_eq6.font.size = Pt(11)
    run_eq6.font.bold = True
    
    add_heading_styled(doc, "8.2.2 Earthing and Protection Systems", 3, theme)
    for item in REPORT_CONTENT["ch8_protection"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    # Chapter 9
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 9\nTechnical Breakdown & Diagnostics of Core Home Appliances", 1, theme)
    
    add_heading_styled(doc, "9.1 Electric Iron", 2, theme)
    add_p(doc, "A. Working Principle", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_iron_principle"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Formula
    p_eq7 = doc.add_paragraph()
    p_eq7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_eq7.paragraph_format.space_before = Pt(4)
    p_eq7.paragraph_format.space_after = Pt(8)
    run_eq7 = p_eq7.add_run("H = I^2Rt     (9.1)")
    run_eq7.font.name = font_name
    run_eq7.font.size = Pt(11)
    run_eq7.font.bold = True
    
    add_p(doc, "B. Key Components and Parts", student_index, font_name, 10.5, bold=True, space_after=3)
    for item in REPORT_CONTENT["ch9_iron_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
    
    add_p(doc, "C. Appliance Conclusion", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_iron_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "9.2 Induction Cooker", 2, theme)
    add_p(doc, "A. Working Principle", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_induction_principle"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "B. Key Components and Parts", student_index, font_name, 10.5, bold=True, space_after=3)
    for item in REPORT_CONTENT["ch9_induction_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_p(doc, "C. Appliance Conclusion", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_induction_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "9.3 Mixer Grinder", 2, theme)
    add_p(doc, "A. Working Principle", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_mixer_principle"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "B. Key Components and Parts", student_index, font_name, 10.5, bold=True, space_after=3)
    for item in REPORT_CONTENT["ch9_mixer_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_heading_styled(doc, "9.3.1 Ceiling Fan", 3, theme)
    add_p(doc, "A. Working Principle", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_ceiling_fan_principle"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "B. Key Components and Parts", student_index, font_name, 10.5, bold=True, space_after=3)
    for item in REPORT_CONTENT["ch9_ceiling_fan_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_p(doc, "C. Appliance Conclusion", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_ceiling_fan_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_styled(doc, "9.3.2 Wall Fan", 3, theme)
    add_p(doc, "A. Working Principle", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_wall_fan_principle"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "B. Key Components and Parts", student_index, font_name, 10.5, bold=True, space_after=3)
    for item in REPORT_CONTENT["ch9_wall_fan_components"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    add_p(doc, "C. Appliance Conclusion", student_index, font_name, 10.5, bold=True, space_after=3)
    add_p(doc, REPORT_CONTENT["ch9_wall_fan_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Chapter 10
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 10\nTroubleshooting & Maintenance Practice Matrix", 1, theme)
    build_styled_table(doc, "troubleshooting_matrix", theme, student_index)
    build_styled_table(doc, "appliance_troubleshooting", theme, student_index)
    
    add_heading_styled(doc, "10.1 Internship Outcomes & Skills Acquired", 2, theme)
    for item in REPORT_CONTENT["ch10_outcomes"]:
        add_list_item(doc, item, student_index, font_name, 10.5)
        
    # Chapter 11
    doc.add_page_break()
    add_heading_styled(doc, "Chapter 11\nApplied Practical Calculations", 1, theme)
    add_heading_styled(doc, "11.1 Domestic Electricity Consumption Auditing", 2, theme)
    add_p(doc, REPORT_CONTENT["ch11_audit_intro"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_p(doc, REPORT_CONTENT["ch11_audit_problem"], student_index, font_name, 11, italic=True)
    
    # Math lines
    p_audit_calc = doc.add_paragraph()
    p_audit_calc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_audit_calc.paragraph_format.left_indent = Inches(0.5)
    p_audit_calc.paragraph_format.line_spacing = 1.25
    p_audit_calc.paragraph_format.space_before = Pt(4)
    p_audit_calc.paragraph_format.space_after = Pt(8)
    run_ac = p_audit_calc.add_run(
        "ELights = 2 × 40 W × 6 hrs/day × 30 days = 14,400 Wh = 14.4 kWh\n"
        "EFans   = 3 × 65 W × 8 hrs/day × 30 days = 46,800 Wh = 46.8 kWh\n"
        "EMisc   = 150 W × 2 hrs/day × 30 days = 9,000 Wh = 9.0 kWh\n"
        "ETotal  = 14.4 kWh + 46.8 kWh + 9.0 kWh = 70.2 kWh (Units)"
    )
    run_ac.font.name = font_name
    run_ac.font.size = Pt(11)
    run_ac.font.bold = True
    
    add_heading_styled(doc, "11.2 Charge Calculation in Audio Systems", 2, theme)
    add_p(doc, REPORT_CONTENT["ch11_charge_problem"], student_index, font_name, 11, italic=True)
    
    p_charge_calc = doc.add_paragraph()
    p_charge_calc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_charge_calc.paragraph_format.left_indent = Inches(0.5)
    p_charge_calc.paragraph_format.line_spacing = 1.25
    p_charge_calc.paragraph_format.space_before = Pt(4)
    p_charge_calc.paragraph_format.space_after = Pt(8)
    run_cc = p_charge_calc.add_run(
        "Given: V = 230 V, E = 8 × 10^6 J\n"
        "Equation: E = qV\n"
        "q = E / V\n"
        "  = 8 × 10^6 J / 230 V\n"
        "  = 3.48 × 10^4 Coulombs (C)"
    )
    run_cc.font.name = font_name
    run_cc.font.size = Pt(11)
    run_cc.font.bold = True
    
    # Part II Conclusion
    add_heading_styled(doc, "Part Conclusion: Domestic Electrical Engineering", 1, theme)
    add_p(doc, REPORT_CONTENT["part2_conclusion"], student_index, font_name, 11, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    # Add footer page numbering
    add_footer_page_number(doc, student_name, theme)
    
    return doc

# Parse pdftotext output to extract page numbers of headings
def parse_pdf_for_page_numbers(pdf_txt_path):
    with open(pdf_txt_path, "r", encoding="utf-8") as f:
        content = f.read()
        
    pages = content.split('\x0c')
    print(f"Temporary PDF parsed. Total text pages: {len(pages)}")
    
    # We want to map each search key to the page it first appears on (excluding the TOC itself, i.e. page 1 to 5)
    page_map = {}
    
    # Keys and regexes
    search_keys = {
        "CERTIFICATE": re.compile(r'CERTIFICATE', re.IGNORECASE),
        "INTERNSHIP ACTIVITY SCHEDULE": re.compile(r'INTERNSHIP\s+ACTIVITY\s+SCHEDULE', re.IGNORECASE),
        "Part I": re.compile(r'Part\s+I', re.IGNORECASE),
        "Part Introduction: Fiber Optic Infrastructure": re.compile(r'Part\s+Introduction:\s*Fiber\s+Optic\s+Infrastructure', re.IGNORECASE),
        "Chapter 1": re.compile(r'Chapter\s+1', re.IGNORECASE),
        "1.1": re.compile(r'1\.1\s+Overview', re.IGNORECASE),
        "1.2": re.compile(r'1\.2\s+The\s+Basic\s+Fiber\s+Optic\s+Link', re.IGNORECASE),
        "1.3": re.compile(r'1\.3\s+Advantages', re.IGNORECASE),
        "Chapter 2": re.compile(r'Chapter\s+2', re.IGNORECASE),
        "2.1": re.compile(r'2\.1\s+Principle\s+of\s+Operation', re.IGNORECASE),
        "2.2": re.compile(r'2\.2\s+Cable\s+Anatomy', re.IGNORECASE),
        "2.3": re.compile(r'2\.3\s+Cable\s+Types', re.IGNORECASE),
        "Chapter 3": re.compile(r'Chapter\s+3', re.IGNORECASE),
        "3.1": re.compile(r'3\.1\s+OFC\s+Splicing', re.IGNORECASE),
        "3.2": re.compile(r'3\.2\s+Logarithmic\s+Conversion', re.IGNORECASE),
        "3.3": re.compile(r'3\.3\s+Splitter\s+Loss', re.IGNORECASE),
        "3.4": re.compile(r'3\.4\s+Standard\s+Loss', re.IGNORECASE),
        "3.5 ": re.compile(r'3\.5\s+Mathematical\s+Verification', re.IGNORECASE),
        "3.5.1": re.compile(r'3\.5\.1\s+Case\s+Study', re.IGNORECASE),
        "3.5.2": re.compile(r'3\.5\.2\s+Link\s+Status', re.IGNORECASE),
        "Chapter 4": re.compile(r'Chapter\s+4', re.IGNORECASE),
        "4.1": re.compile(r'4\.1\s+What\s+is\s+Splicing', re.IGNORECASE),
        "4.2": re.compile(r'4\.2\s+Types\s+of\s+Splicing', re.IGNORECASE),
        "4.2.1": re.compile(r'4\.2\.1\s+Fusion\s+Splicing', re.IGNORECASE),
        "4.2.2": re.compile(r'4\.2\.2\s+Mechanical\s+Splicing', re.IGNORECASE),
        "4.3": re.compile(r'4\.3\s+Importance', re.IGNORECASE),
        "4.4": re.compile(r'4\.4\s+Step-by-Step', re.IGNORECASE),
        "4.5": re.compile(r'4\.5\s+Advantages', re.IGNORECASE),
        "4.6": re.compile(r'4\.6\s+Future\s+Scope', re.IGNORECASE),
        "4.7": re.compile(r'4\.7\s+Conclusion', re.IGNORECASE),
        "Chapter 5": re.compile(r'Chapter\s+5', re.IGNORECASE),
        "5.1": re.compile(r'5\.1\s+Overview\s+of\s+FTTH', re.IGNORECASE),
        "5.2": re.compile(r'5\.2\s+The\s+FTTH\s+Network', re.IGNORECASE),
        "5.2.1": re.compile(r'A\.\s+The\s+Feeder\s+Segment', re.IGNORECASE),
        "5.2.2": re.compile(r'B\.\s+The\s+Distribution\s+Segment', re.IGNORECASE),
        "5.2.3": re.compile(r'C\.\s+The\s+Drop\s+Segment', re.IGNORECASE),
        "5.3": re.compile(r'5\.3\s+Core\s+Hardware', re.IGNORECASE),
        "5.3.1": re.compile(r'5\.3\.1\s+Active\s+Components', re.IGNORECASE),
        "5.3.2": re.compile(r'5\.3\.2\s+Passive\s+Components', re.IGNORECASE),
        "Chapter 6": re.compile(r'Chapter\s+6', re.IGNORECASE),
        "6.1": re.compile(r'6\.1\s+Fiber\s+Color\s+Coding', re.IGNORECASE),
        "6.2": re.compile(r'6\.2\s+Physical\s+Deployment', re.IGNORECASE),
        "6.2.1": re.compile(r'6\.2\.1\s+Site\s+Survey', re.IGNORECASE),
        "6.2.2": re.compile(r'6\.2\.2\s+Outdoor\s+Routing', re.IGNORECASE),
        "6.2.3": re.compile(r'6\.2\.3\s+Indoor\s+Termination', re.IGNORECASE),
        "6.3": re.compile(r'6\.3\s+Splicing\s+Methods', re.IGNORECASE),
        "6.4": re.compile(r'6\.4\s+Link\s+Budget', re.IGNORECASE),
        "6.4.1": re.compile(r'6\.4\.1\s+Wavelength', re.IGNORECASE),
        "6.4.2": re.compile(r'6\.4\.2\s+Testing\s+Parameters', re.IGNORECASE),
        "6.5": re.compile(r'6\.5\s+Troubleshooting', re.IGNORECASE),
        "Chapter 7": re.compile(r'Chapter\s+7', re.IGNORECASE),
        "7.1": re.compile(r'7\.1\s+Introduction\s+to\s+the\s+Field', re.IGNORECASE),
        "7.2": re.compile(r'7\.2\s+Observations\s+and\s+Core', re.IGNORECASE),
        "7.3": re.compile(r'7\.3\s+Conclusion\s+of\s+the\s+Visit', re.IGNORECASE),
        "Part Conclusion: Fiber Optic Infrastructure": re.compile(r'Part\s+Conclusion:\s*Fiber\s+Optic\s+Infrastructure', re.IGNORECASE),
        "Part II": re.compile(r'Part\s+II', re.IGNORECASE),
        "Part Introduction: Domestic Electrical Engineering": re.compile(r'Part\s+Introduction:\s*Domestic\s+Electrical\s+Engineering', re.IGNORECASE),
        "Chapter 8": re.compile(r'Chapter\s+8', re.IGNORECASE),
        "8.1": re.compile(r'8\.1\s+Introduction', re.IGNORECASE),
        "8.2": re.compile(r'8\.2\s+Fundamentals\s+of\s+Electrical', re.IGNORECASE),
        "8.2.1": re.compile(r'8\.2\.1\s+Circuit\s+Behaviors', re.IGNORECASE),
        "8.2.2": re.compile(r'8\.2\.2\s+Earthing\s+and\s+Protection', re.IGNORECASE),
        "Chapter 9": re.compile(r'Chapter\s+9', re.IGNORECASE),
        "9.1": re.compile(r'9\.1\s+Electric\s+Iron', re.IGNORECASE),
        "9.2": re.compile(r'9\.2\s+Induction\s+Cooker', re.IGNORECASE),
        "9.3 ": re.compile(r'9\.3\s+Mixer\s+Grinder', re.IGNORECASE),
        "9.3.1": re.compile(r'9\.3\.1\s+Ceiling\s+Fan', re.IGNORECASE),
        "9.3.2": re.compile(r'9\.3\.2\s+Wall\s+Fan', re.IGNORECASE),
        "Chapter 10": re.compile(r'Chapter\s+10', re.IGNORECASE),
        "10.1": re.compile(r'10\.1\s+Internship Outcomes', re.IGNORECASE),
        "Chapter 11": re.compile(r'Chapter 11', re.IGNORECASE),
        "11.1": re.compile(r'11\.1\s+Domestic Electricity', re.IGNORECASE),
        "11.2": re.compile(r'11\.2\s+Charge Calculation', re.IGNORECASE),
        "Part Conclusion: Domestic Electrical Engineering": re.compile(r'Part Conclusion:\s*Domestic Electrical', re.IGNORECASE)
    }
    
    # Hardcoded known positions for the early pages
    page_map["CERTIFICATE"] = 2
    
    # For each page from 2 to the end, search for keys
    for p_idx in range(1, len(pages)): # 1-indexed (pages[0] is Page 1)
        page_num = p_idx + 1
        page_text = pages[p_idx]
        
        # Check if this looks like a TOC page (if it contains too many dots or the word "CONTENTS")
        is_toc_page = "CONTENTS" in page_text or "Table of Contents" in page_text or page_text.count("...") > 10
        if is_toc_page and page_num > 2:
            continue # Skip finding headings in the TOC itself!
            
        for key, regex in search_keys.items():
            # If we haven't mapped this key yet, and it matches on this page
            if key not in page_map and regex.search(page_text):
                # Ensure we don't map Part I / Part II / Chapters to early pages
                if key.startswith("Chapter") or key.startswith("Part") or "." in key:
                    if page_num <= 4:
                        continue
                page_map[key] = page_num
                
    # Double check if any key was missed and default it
    for key in search_keys:
        if key not in page_map:
            print(f"Warning: could not locate page for key '{key}' in PDF text. Defaulting.")
            page_map[key] = "XX"
            
    return page_map

# Compile single student report
def generate_student_report(student_name, student_index):
    theme = THEMES[student_name]
    safe_name = student_name.replace(" ", "_")
    
    print(f"\n==================================================")
    print(f"Processing Student {student_index + 1}/7: {student_name}")
    print(f"Theme Font: {theme['font']}, Cover: {theme['cover_style']}")
    print(f"==================================================")
    
    # Pass 1: Build docx with placeholder page numbers
    print("Pass 1: Creating temporary document...")
    doc_temp = build_report_docx(student_name, theme, student_index, page_numbers=None)
    temp_docx_path = os.path.join(OUTPUT_DIR, f"{safe_name}_temp.docx")
    doc_temp.save(temp_docx_path)
    
    # Convert to PDF using LibreOffice
    print("Pass 1: Converting to temporary PDF...")
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", OUTPUT_DIR, temp_docx_path]
    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    
    temp_pdf_path = os.path.join(OUTPUT_DIR, f"{safe_name}_temp.pdf")
    temp_txt_path = os.path.join(OUTPUT_DIR, f"{safe_name}_temp_text.txt")
    
    if not os.path.exists(temp_pdf_path):
        print(f"Error: Temporary PDF for {student_name} was not generated.")
        return False
        
    # Convert PDF to text for analysis
    subprocess.run(["pdftotext", temp_pdf_path, temp_txt_path])
    
    # Parse text to find page numbers
    page_numbers = parse_pdf_for_page_numbers(temp_txt_path)
    print(f"Extracted Page Map: {page_numbers}")
    
    # Pass 2: Rebuild docx with correct page numbers
    print("Pass 2: Creating final document with updated TOC page numbers...")
    doc_final = build_report_docx(student_name, theme, student_index, page_numbers)
    final_docx_path = os.path.join(OUTPUT_DIR, f"{safe_name}_Internship_Report.docx")
    doc_final.save(final_docx_path)
    
    # Convert to final PDF
    print("Pass 2: Converting to final PDF...")
    cmd_final = ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", OUTPUT_DIR, final_docx_path]
    subprocess.run(cmd_final, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    
    final_pdf_path = os.path.join(OUTPUT_DIR, f"{safe_name}_Internship_Report.pdf")
    if os.path.exists(final_pdf_path):
        print(f"Success! Generated {final_pdf_path} (size: {os.path.getsize(final_pdf_path)} bytes)")
    else:
        print(f"Error: Final PDF for {student_name} was not generated.")
        return False
        
    # Clean up temp files
    for path in [temp_docx_path, temp_pdf_path, temp_txt_path]:
        if os.path.exists(path):
            os.remove(path)
            
    return True

# Main Execution Loop
def main():
    print("Starting internship report generation for physics students...")
    success_count = 0
    
    for idx, name in enumerate(STUDENTS):
        success = generate_student_report(name, idx)
        if success:
            success_count += 1
            
    print("\n==================================================")
    print(f"Generation complete. Successfully generated {success_count}/7 reports.")
    print("==================================================")

if __name__ == "__main__":
    main()
